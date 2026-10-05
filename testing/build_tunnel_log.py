from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import FormulaRule

N_RUNS, N_COMB, N_DES = 40, 20, 10
R0 = 5                                   # first data row on every table
F = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
GRP_FILL = PatternFill("solid", fgColor="D9E1F2")
YEL = PatternFill("solid", fgColor="FFFF00")
BLUE, BLACK, GREEN = "0000FF", "000000", "008000"
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

def font(color=BLACK, bold=False, size=10, italic=False):
    return Font(name=F, color=color, bold=bold, size=size, italic=italic)

def title(ws, text, sub):
    ws["A1"] = text; ws["A1"].font = font(bold=True, size=14)
    ws["A2"] = sub; ws["A2"].font = font(italic=True, size=9, color="595959")

def headers(ws, row, names, widths):
    for i, (n, w) in enumerate(zip(names, widths), 1):
        c = ws.cell(row=row, column=i, value=n)
        c.font = Font(name=F, bold=True, color="FFFFFF", size=10)
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[row].height = 42
    ws.freeze_panes = ws.cell(row=row + 1, column=1)

wb = Workbook()

# ---------------- Settings ----------------
st = wb.active; st.title = "Settings"
title(st, "Air brake wind tunnel test log",
      "Log every tunnel run on the Tunnel runs tab. Combined and Summary update by themselves.")
st["A4"] = "How to use this workbook"; st["A4"].font = font(bold=True, size=11)
howto = [
    "1. Add each air brake version on the Designs tab (one row per version).",
    "2. Log every tunnel measurement on Tunnel runs: one row per section (Nose, Air brake or Fin can) per run.",
    "3. On Combined, add a row for each design, speed setpoint and deployment you tested. It averages repeat runs and adds the three sections together.",
    "4. Summary shows drag coefficient against deployment and an estimated apogee for each design.",
    "Blue text: values you type.  Yellow fill: needs your number.  Black text: formulas, leave them alone.",
    "Need more rows? Select the last filled row on Tunnel runs or Combined and drag it down; the formulas copy with it.",
]
for i, t in enumerate(howto):
    st.cell(row=5 + i, column=1, value=t).font = font()
st["A12"] = "Constants and inputs"; st["A12"].font = font(bold=True, size=11)
headers_row = 13
for i, h in enumerate(["Setting", "Value", "Unit", "Source / note"], 1):
    c = st.cell(row=headers_row, column=i, value=h)
    c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = HDR_FILL; c.border = BOX
settings = [
    ("Gas constant for dry air", 287.05, "J/(kg·K)", "Standard value for dry air", False),
    ("Sutherland viscosity constant C1", 1.458e-6, "kg/(m·s·K^0.5)", "Sutherland's law for air (standard constants)", False),
    ("Sutherland temperature S", 110.4, "K", "Sutherland's law for air (standard constants)", False),
    ("Gravity", 9.81, "m/s²", "Standard gravity", False),
    ("Feet per metre", 3.28084, "ft/m", "Unit conversion", False),
    ("Air density at launch site", 1.225, "kg/m³", "ISA sea level. Replace with your launch site value.", True),
    ("Target apogee", 800, "ft", "Source: user-provided (ARC target)", True),
    ("Tunnel test-section area", None, "m²", "Fill in: width × height of your tunnel's test section", True),
    ("Blockage limit", 0.10, "fraction", "Rule of thumb: above about 10% the walls inflate the drag reading", True),
]
for i, (name, val, unit, note, user) in enumerate(settings):
    r = headers_row + 1 + i
    st.cell(row=r, column=1, value=name).font = font()
    v = st.cell(row=r, column=2, value=val); v.font = font(BLUE)
    if user: v.fill = YEL
    st.cell(row=r, column=3, value=unit).font = font()
    st.cell(row=r, column=4, value=note).font = font(size=9, color="595959")
    for col in range(1, 5): st.cell(row=r, column=col).border = BOX
st["B22"].number_format = "0%"
st["B15"].number_format = "0.000E+00"
for col, w in zip("ABCD", (34, 14, 16, 70)):
    st.column_dimensions[col].width = w
S = {k: "Settings!$B$%d" % (14 + i) for i, k in enumerate(
    ["R", "C1", "S", "g", "ftm", "rho", "target", "tunnel", "blk"])}

# ---------------- Designs ----------------
de = wb.create_sheet("Designs")
title(de, "Air brake designs", "One row per version. Design ID is what the other tabs look up, so keep it short and unique.")
dcols = ["Design ID", "Description", "Body tube OD (mm)", "Reference area (m²)", "Number of leaves",
         "Leaf width (mm)", "Leaf extension (mm)", "Brake area, fully open (cm²)", "Servo",
         "Rocket mass (g)", "CAD file / link", "Notes"]
headers(de, 4, dcols, (14, 30, 12, 12, 10, 11, 11, 13, 22, 11, 30, 40))
for r in range(R0, R0 + N_DES):
    de.cell(row=r, column=4, value=f'=IF(C{r}="","",PI()*(C{r}/2000)^2)')
    de.cell(row=r, column=8, value=f'=IF(OR(E{r}="",F{r}="",G{r}=""),"",E{r}*F{r}*G{r}/100)')
    for c in range(1, 13):
        cell = de.cell(row=r, column=c); cell.border = BOX
        cell.font = font(BLACK if c in (4, 8) else BLUE)
    de.cell(row=r, column=4).number_format = "0.00000"
    de.cell(row=r, column=8).number_format = "0.0"
ex = ["MIT-6in", "Example: MIT Rocket Team CDR design (replace with your own)", 152.4, None, 4, 88.9, 57.15,
      None, "1500 oz-in industrial servo", None, "airbrake.fs", "Leaf size from the MIT wiki"]
for c, v in enumerate(ex, 1):
    if v is not None: de.cell(row=R0, column=c, value=v)
de.cell(row=R0, column=1).comment = Comment("Example row showing the format. Source: MIT Rocket Team wiki, Air Brakes page. Replace or delete.", "Workbook")

# ---------------- Tunnel runs ----------------
ru = wb.create_sheet("Tunnel runs")
title(ru, "Wind tunnel runs", "One row per section per run. Grey columns are calculated.")
rcols = ["Run ID", "Date", "Design ID", "Section", "Deployment (%)", "Speed setpoint (m/s)",
         "Measured airspeed (m/s)", "Measured force (N)", "Tare force (N)", "Air temp (°C)",
         "Air pressure (kPa)", "Model frontal area (cm²)", "Notes",
         "Air density (kg/m³)", "Dynamic pressure (Pa)", "Net drag (N)", "Body OD (m)",
         "Reference area (m²)", "Section drag coefficient", "Air viscosity (Pa·s)",
         "Reynolds number", "Blockage", "Blockage check", "Condition key"]
headers(ru, 4, rcols, (10, 11, 12, 11, 11, 11, 11, 11, 10, 9, 10, 11, 28,
                       10, 10, 10, 9, 11, 11, 12, 12, 9, 10, 26))
calc_fill = PatternFill("solid", fgColor="F2F2F2")
for r in range(R0, R0 + N_RUNS):
    f = {
        14: f'=IF($A{r}="","",$K{r}*1000/({S["R"]}*($J{r}+273.15)))',
        15: f'=IF($A{r}="","",0.5*N{r}*G{r}^2)',
        16: f'=IF($A{r}="","",H{r}-I{r})',
        17: f'=IF(COUNTIF(Designs!$A:$A,$C{r})=0,"",INDEX(Designs!$C:$C,MATCH($C{r},Designs!$A:$A,0))/1000)',
        18: f'=IF(COUNTIF(Designs!$A:$A,$C{r})=0,"",INDEX(Designs!$D:$D,MATCH($C{r},Designs!$A:$A,0)))',
        19: f'=IF(OR($A{r}="",O{r}="",R{r}=""),"",IF(O{r}*R{r}=0,"",P{r}/(O{r}*R{r})))',
        20: f'=IF($A{r}="","",{S["C1"]}*($J{r}+273.15)^1.5/($J{r}+273.15+{S["S"]}))',
        21: f'=IF(OR($A{r}="",Q{r}=""),"",N{r}*G{r}*Q{r}/T{r})',
        22: f'=IF(OR(L{r}="",{S["tunnel"]}=""),"",L{r}/10000/{S["tunnel"]})',
        23: f'=IF(V{r}="","",IF(V{r}>{S["blk"]},"Too high","OK"))',
        24: f'=IF($A{r}="","",$C{r}&"|"&$F{r}&"|"&$E{r}&"|"&$D{r})',
    }
    for c in range(1, 25):
        cell = ru.cell(row=r, column=c)
        if c in f:
            cell.value = f[c]; cell.font = font(GREEN if c in (17, 18) else BLACK); cell.fill = calc_fill
    ru.cell(row=r, column=2).number_format = "yyyy-mm-dd"
    for c, fmt in ((14, "0.000"), (15, "0.0"), (16, "0.000"), (17, "0.000"), (18, "0.00000"),
                   (19, "0.000"), (20, "0.00E+00"), (21, "#,##0"), (22, "0.0%")):
        ru.cell(row=r, column=c).number_format = fmt
example_run = ["EX-1", None, "MIT-6in", "Air brake", 100, 20, 20.1, 6.40, 0.30, 21, 101.3, 182.4,
               "Example row showing the format. Delete it."]
for c, v in enumerate(example_run, 1):
    if v is not None: ru.cell(row=R0, column=c, value=v)
from datetime import date
ru.cell(row=R0, column=2, value=date(2026, 10, 4))
ru["H4"].comment = Comment("Force from the balance, in newtons. If your scale reads grams, multiply by 0.00981.", "Workbook")
ru["I4"].comment = Comment("Force with only the mount/sting in the tunnel at the same speed. It is subtracted from the measured force.", "Workbook")
ru["X4"].comment = Comment("Design | speed setpoint | deployment | section. Combined matches runs on this.", "Workbook")
ru["F4"].comment = Comment("The nominal tunnel setting. Combined matches runs on this value, so type it the same way every time (for example 20, not 20.1).", "Workbook")

dv_design = DataValidation(type="list", formula1=f"=Designs!$A${R0}:$A${R0+N_DES-1}", allow_blank=True)
dv_sec = DataValidation(type="list", formula1='"Nose,Air brake,Fin can"', allow_blank=True)
dv_dep = DataValidation(type="list", formula1='"0,25,50,75,100"', allow_blank=True)
for dv in (dv_design, dv_sec, dv_dep): ru.add_data_validation(dv)
dv_design.add(f"C{R0}:C{R0+N_RUNS-1}"); dv_sec.add(f"D{R0}:D{R0+N_RUNS-1}"); dv_dep.add(f"E{R0}:E{R0+N_RUNS-1}")
ru.conditional_formatting.add(f"W{R0}:W{R0+N_RUNS-1}",
    FormulaRule(formula=[f'$W{R0}="Too high"'], fill=PatternFill("solid", fgColor="F4CCCC")))

# ---------------- Combined ----------------
co = wb.create_sheet("Combined")
title(co, "Whole-rocket drag from the three sections",
      "Add one row per design, speed setpoint and deployment. Repeat runs are averaged; the three sections are added together.")
ccols = ["Design ID", "Speed setpoint (m/s)", "Deployment (%)", "Nose", "Air brake", "Fin can",
         "All three sections tested?", "Total drag (N)", "Average dynamic pressure (Pa)",
         "Reference area (m²)", "Whole-rocket drag coefficient", "Air brake section drag, stowed (N)",
         "Drag added by brakes (N)", "Drag coefficient added by brakes"]
headers(co, 4, ccols, (12, 11, 11, 10, 10, 10, 12, 11, 12, 11, 13, 13, 12, 13))
co.merge_cells("D3:F3"); co["D3"] = "Average net drag by section (N)"
co["D3"].font = font(bold=True); co["D3"].fill = GRP_FILL; co["D3"].alignment = Alignment(horizontal="center")
RU = "'Tunnel runs'!"
for r in range(R0, R0 + N_COMB):
    f = {}
    key = f'$A{r}&"|"&$B{r}&"|"&$C{r}&"|"'
    for c, col in ((4, "D"), (5, "E"), (6, "F")):
        f[c] = (f'=IF($A{r}="","",IF(COUNTIF({RU}$X:$X,{key}&{col}$4)=0,"",'
                f'AVERAGEIF({RU}$X:$X,{key}&{col}$4,{RU}$P:$P)))')
    f[7] = f'=IF($A{r}="","",IF(COUNT(D{r}:F{r})=3,"Yes","No"))'
    f[8] = f'=IF(G{r}="Yes",SUM(D{r}:F{r}),"")'
    pre = f'$A{r}&"|"&$B{r}&"|"&$C{r}&"|*"'
    f[9] = f'=IF($A{r}="","",IF(COUNTIF({RU}$X:$X,{pre})=0,"",AVERAGEIF({RU}$X:$X,{pre},{RU}$O:$O)))'
    f[10] = f'=IF(COUNTIF(Designs!$A:$A,$A{r})=0,"",INDEX(Designs!$D:$D,MATCH($A{r},Designs!$A:$A,0)))'
    f[11] = f'=IF(OR(H{r}="",I{r}="",J{r}=""),"",H{r}/(I{r}*J{r}))'
    stow = f'$A{r}&"|"&$B{r}&"|0|"&E$4'
    f[12] = f'=IF($A{r}="","",IF(COUNTIF({RU}$X:$X,{stow})=0,"",AVERAGEIF({RU}$X:$X,{stow},{RU}$P:$P)))'
    f[13] = f'=IF(OR(E{r}="",L{r}=""),"",E{r}-L{r})'
    f[14] = f'=IF(OR(M{r}="",I{r}="",J{r}=""),"",M{r}/(I{r}*J{r}))'
    for c in range(1, 15):
        cell = co.cell(row=r, column=c)
        if c in f:
            cell.value = f[c]; cell.font = font(GREEN if c == 10 else BLACK); cell.fill = calc_fill
    for c, fmt in ((4, "0.000"), (5, "0.000"), (6, "0.000"), (8, "0.000"), (9, "0.0"),
                   (10, "0.00000"), (11, "0.000"), (12, "0.000"), (13, "0.000"), (14, "0.000")):
        co.cell(row=r, column=c).number_format = fmt
for c, v in enumerate(["MIT-6in", 20, 100], 1):
    co.cell(row=R0, column=c, value=v)
co["L4"].comment = Comment("The air brake section at 0% deployment, same design and speed. Subtracting it cancels the extra drag from the cut ends.", "Workbook")
dvc_design = DataValidation(type="list", formula1=f"=Designs!$A${R0}:$A${R0+N_DES-1}", allow_blank=True)
dvc_dep = DataValidation(type="list", formula1='"0,25,50,75,100"', allow_blank=True)
co.add_data_validation(dvc_design); co.add_data_validation(dvc_dep)
dvc_design.add(f"A{R0}:A{R0+N_COMB-1}"); dvc_dep.add(f"C{R0}:C{R0+N_COMB-1}")
co.conditional_formatting.add(f"G{R0}:G{R0+N_COMB-1}",
    FormulaRule(formula=[f'$G{R0}="No"'], fill=PatternFill("solid", fgColor="FCE4D6")))

# ---------------- Summary ----------------
su = wb.create_sheet("Summary")
title(su, "Summary by design", "Drag coefficients are averaged over every speed with all three sections tested.")
su["A4"] = "Whole-rocket drag coefficient by brake deployment"; su["A4"].font = font(bold=True, size=11)
deps = [0, 25, 50, 75, 100]
shdr = ["Design ID"] + [f"{d}" for d in deps] + ["Brake drag coefficient added at 100%", "Tunnel runs logged"]
headers(su, 5, shdr, (14, 11, 11, 11, 11, 11, 16, 12))
su.freeze_panes = None
for i, d in enumerate(deps):
    su.cell(row=5, column=2 + i, value=d).number_format = '0"% open"'
CO = "Combined!"
g0 = 6
for k in range(N_DES):
    r = g0 + k
    su.cell(row=r, column=1, value=f'=IF(Designs!$A{R0+k}="","",Designs!$A{R0+k})').font = font(GREEN)
    for i in range(5):
        col = su.cell(row=5, column=2 + i).column_letter
        cr = f'{CO}$A:$A,$A{r},{CO}$C:$C,{col}$5,{CO}$G:$G,"Yes"'
        su.cell(row=r, column=2 + i,
                value=f'=IF($A{r}="","",IF(COUNTIFS({cr})=0,"",AVERAGEIFS({CO}$K:$K,{cr})))').number_format = "0.000"
    cr = f'{CO}$A:$A,$A{r},{CO}$C:$C,100'
    su.cell(row=r, column=7, value=(f'=IF($A{r}="","",IF(COUNTIFS({cr},{CO}$N:$N,">-1E+300")=0,"",'
                                    f'AVERAGEIFS({CO}$N:$N,{cr},{CO}$N:$N,">-1E+300")))')).number_format = "0.000"
    su.cell(row=r, column=8, value=f'=IF($A{r}="","",COUNTIF(\'Tunnel runs\'!$C:$C,$A{r}))')
    for c in range(1, 9):
        su.cell(row=r, column=c).border = BOX
        if c > 1: su.cell(row=r, column=c).font = font()

a0 = g0 + N_DES + 2
su.cell(row=a0, column=1, value="Estimated apogee").font = font(bold=True, size=11)
su.cell(row=a0 + 1, column=1, value=("Vertical coast with constant drag, from burnout to apogee. Take burnout altitude and "
     "velocity from OpenRocket. Tunnel speeds are lower than flight speeds, so treat this as a comparison between designs.")).font = font(size=9, italic=True, color="595959")
ahdr = ["Design ID", "Coast mass (kg)", "Burnout altitude (m)", "Burnout velocity (m/s)",
        "Drag coefficient, stowed", "Drag coefficient, fully open", "Reference area (m²)",
        "Apogee, brakes stowed (ft)", "Apogee, brakes open all coast (ft)", "Target apogee (ft)", "Can the brakes reach the target?"]
headers(su, a0 + 2, ahdr, (14, 11, 11, 11, 12, 12, 12, 13, 14, 11, 22))
su.freeze_panes = None
for k in range(N_DES):
    r = a0 + 3 + k
    gr = g0 + k
    apo = lambda cd: (f'=IF(OR($A{r}="",$B{r}="",$C{r}="",$D{r}="",{cd}{r}="",$G{r}=""),"",'
                      f'($C{r}+$B{r}/({S["rho"]}*{cd}{r}*$G{r})*LN(1+{S["rho"]}*{cd}{r}*$G{r}*$D{r}^2/(2*$B{r}*{S["g"]})))*{S["ftm"]})')
    vals = {1: f'=$A{gr}', 5: f'=IF($A{r}="","",B{gr})', 6: f'=IF($A{r}="","",F{gr})',
            7: f'=IF(COUNTIF(Designs!$A:$A,$A{r})=0,"",INDEX(Designs!$D:$D,MATCH($A{r},Designs!$A:$A,0)))',
            8: apo("E"), 9: apo("F"), 10: f'=IF($A{r}="","",{S["target"]})',
            11: (f'=IF(OR(H{r}="",I{r}=""),"",IF(H{r}<J{r},"No: too low even stowed",'
                 f'IF(I{r}>J{r},"No: too high even fully open","Yes")))')}
    for c in range(1, 12):
        cell = su.cell(row=r, column=c); cell.border = BOX
        if c in vals:
            cell.value = vals[c]; cell.font = font(GREEN if c in (1, 7, 10) else BLACK)
        else:
            cell.font = font(BLUE); cell.fill = YEL
    for c, fmt in ((5, "0.000"), (6, "0.000"), (7, "0.00000"), (8, "#,##0"), (9, "#,##0"), (10, "#,##0")):
        su.cell(row=r, column=c).number_format = fmt
su.column_dimensions["K"].width = 26

ch = LineChart()
ch.title = "Whole-rocket drag coefficient vs brake deployment"
ch.y_axis.title = "Drag coefficient"; ch.x_axis.title = "Deployment (%)"
ch.height, ch.width = 8, 16
data = Reference(su, min_col=1, max_col=6, min_row=g0, max_row=g0 + N_DES - 1)
ch.add_data(data, from_rows=True, titles_from_data=True)
ch.set_categories(Reference(su, min_col=2, max_col=6, min_row=5, max_row=5))
su.add_chart(ch, "M4")

for ws in (ru, co, de):
    for col in range(1, 25):
        ws.column_dimensions[ws.cell(row=4, column=col).column_letter].font = font(BLUE)
wb.save("airbrake_tunnel_log.xlsx")
print("saved")
